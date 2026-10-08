import sys
import tempfile
import unittest
from pathlib import Path


SPARK_DIR = Path(__file__).resolve().parent.parent
if str(SPARK_DIR) not in sys.path:
    sys.path.insert(0, str(SPARK_DIR))

from ai_dataset_router import _extract_docx_text


class DocumentRouterTest(unittest.TestCase):
    def test_docx_is_extracted_as_text_with_tables(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("Báo cáo học tập")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "ma_sv"
            table.cell(0, 1).text = "diem"
            table.cell(1, 0).text = "SV001"
            table.cell(1, 1).text = "9.0"
            document.save(path)

            extracted = _extract_docx_text(path)

        self.assertIn("Báo cáo học tập", extracted)
        self.assertIn("[BẢNG 1]", extracted)
        self.assertIn("ma_sv\tdiem", extracted)
        self.assertIn("SV001\t9.0", extracted)


if __name__ == "__main__":
    unittest.main(verbosity=2)
