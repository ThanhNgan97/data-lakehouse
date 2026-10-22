"""Read-only validation for one object uploaded through Apache Ozone S3 Gateway."""

from __future__ import annotations

import csv
import hashlib
import json
import os
from datetime import datetime, timedelta
from pathlib import Path

import boto3
from airflow import DAG
from airflow.operators.python import PythonOperator
from botocore.config import Config


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".json",
    ".parquet",
    ".xlsx",
    ".xls",
    ".pdf",
    ".docx",
}


def _s3_client():
    return boto3.client(
        "s3",
        endpoint_url=os.environ["S3_ENDPOINT"],
        aws_access_key_id=os.environ["S3_ACCESS_KEY"],
        aws_secret_access_key=os.environ["S3_SECRET_KEY"],
        region_name=os.environ.get("S3_REGION", "us-east-1"),
        config=Config(s3={"addressing_style": "path"}),
    )


def validate_request(**context):
    conf = dict((context.get("dag_run") and context["dag_run"].conf) or {})
    required = ("upload_id", "bucket", "object_key", "expected_size", "expected_sha256")
    missing = [name for name in required if conf.get(name) in (None, "")]
    if missing:
        raise ValueError(f"Missing required DAG configuration: {', '.join(missing)}")

    expected_bucket = os.environ.get("S3_BUCKET_NAME", "university-lakehouse")
    if conf["bucket"] != expected_bucket:
        raise ValueError("The requested bucket is not the configured Ozone bucket.")

    staging_prefix = os.environ.get("STAGING_PREFIX", "staging").strip("/") + "/"
    object_key = str(conf["object_key"]).lstrip("/")
    if not object_key.startswith(staging_prefix) or ".." in Path(object_key).parts:
        raise ValueError("The object key must be located below the configured staging prefix.")

    extension = Path(object_key).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: {extension}")

    expected_sha256 = str(conf["expected_sha256"]).lower()
    if len(expected_sha256) != 64 or any(c not in "0123456789abcdef" for c in expected_sha256):
        raise ValueError("expected_sha256 must be a hexadecimal SHA-256 digest.")

    request = {
        "upload_id": str(conf["upload_id"]),
        "bucket": expected_bucket,
        "object_key": object_key,
        "filename": str(conf.get("filename") or Path(object_key).name),
        "extension": extension,
        "expected_size": int(conf["expected_size"]),
        "expected_sha256": expected_sha256,
    }
    context["ti"].xcom_push(key="request", value=request)
    return request


def head_ozone_object(**context):
    request = context["ti"].xcom_pull(task_ids="validate_request", key="request")
    response = _s3_client().head_object(
        Bucket=request["bucket"],
        Key=request["object_key"],
    )
    actual_size = int(response["ContentLength"])
    if actual_size <= 0 or actual_size != request["expected_size"]:
        raise ValueError(
            f"Ozone object size mismatch: expected {request['expected_size']}, got {actual_size}."
        )
    return {"content_length": actual_size, "etag": response.get("ETag", "").strip('"')}


def download_and_verify_checksum(**context):
    ti = context["ti"]
    request = ti.xcom_pull(task_ids="validate_request", key="request")
    run_id = "".join(c if c.isalnum() or c in "-_" else "_" for c in context["run_id"])
    target_dir = Path("/tmp/ozone-validation") / run_id
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / Path(request["object_key"]).name
    _s3_client().download_file(request["bucket"], request["object_key"], str(target))

    digest = hashlib.sha256()
    with target.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    actual_sha256 = digest.hexdigest()
    if actual_sha256 != request["expected_sha256"]:
        raise ValueError("The downloaded Ozone object does not match the upload checksum.")

    result = {"local_path": str(target), "actual_sha256": actual_sha256}
    ti.xcom_push(key="download", value=result)
    return result


def read_file_sample(**context):
    ti = context["ti"]
    request = ti.xcom_pull(task_ids="validate_request", key="request")
    download = ti.xcom_pull(task_ids="download_and_verify_checksum", key="download")
    path = Path(download["local_path"])
    extension = request["extension"]
    result = {"format": extension.lstrip("."), "readable": True}

    if extension in {".csv", ".tsv"}:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter="\t" if extension == ".tsv" else ",")
            rows = []
            for _, row in zip(range(11), reader):
                rows.append(row)
        result.update(
            columns=(rows[0] if rows else []),
            column_count=(len(rows[0]) if rows else 0),
            sample_row_count=max(0, len(rows) - 1),
        )
    elif extension == ".json":
        with path.open("r", encoding="utf-8-sig") as handle:
            payload = json.load(handle)
        rows = payload if isinstance(payload, list) else [payload]
        first = rows[0] if rows else {}
        result.update(
            columns=(list(first) if isinstance(first, dict) else []),
            sample_row_count=min(10, len(rows)),
        )
    elif extension == ".parquet":
        import pandas as pd

        frame = pd.read_parquet(path).head(10)
        result.update(columns=list(frame.columns), sample_row_count=len(frame))
    elif extension in {".xlsx", ".xls"}:
        import pandas as pd

        frame = pd.read_excel(path, nrows=10)
        result.update(columns=list(frame.columns), sample_row_count=len(frame))
    elif extension == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            result.update(page_count=len(pdf.pages), sample_row_count=min(1, len(pdf.pages)))
    elif extension == ".docx":
        from docx import Document

        document = Document(path)
        result.update(
            paragraph_count=len(document.paragraphs),
            table_count=len(document.tables),
            sample_row_count=min(10, len(document.paragraphs)),
        )

    ti.xcom_push(key="sample", value=result)
    return result


def publish_validation_result(**context):
    ti = context["ti"]
    request = ti.xcom_pull(task_ids="validate_request", key="request")
    sample = ti.xcom_pull(task_ids="read_file_sample", key="sample")
    return {
        "status": "PASS",
        "storage": "apache_ozone",
        "endpoint": os.environ["S3_ENDPOINT"],
        "bucket": request["bucket"],
        "object_key": request["object_key"],
        "upload_id": request["upload_id"],
        "checksum_valid": True,
        "sample": sample,
    }


with DAG(
    dag_id="ozone_read_validation",
    description="Validate that Airflow can read one exact uploaded object from Apache Ozone",
    start_date=datetime(2026, 1, 1),
    schedule_interval=None,
    catchup=False,
    default_args={"owner": "lakehouse", "retries": 1, "retry_delay": timedelta(seconds=10)},
    tags=["ozone", "storage", "validation"],
) as dag:
    validate = PythonOperator(task_id="validate_request", python_callable=validate_request)
    head = PythonOperator(task_id="head_ozone_object", python_callable=head_ozone_object)
    download = PythonOperator(
        task_id="download_and_verify_checksum",
        python_callable=download_and_verify_checksum,
    )
    sample = PythonOperator(task_id="read_file_sample", python_callable=read_file_sample)
    publish = PythonOperator(
        task_id="publish_validation_result",
        python_callable=publish_validation_result,
    )

    validate >> head >> download >> sample >> publish
